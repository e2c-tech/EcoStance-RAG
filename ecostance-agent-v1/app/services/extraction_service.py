import os
import anyio
import fitz  # PyMuPDF, installed as PyMuPDF
import docx
import pandas as pd
from trafilatura import extract as trafilatura_extract
import json
import re
from typing import List, Dict, Any, Tuple, Optional
from bs4 import BeautifulSoup
import sqlparse
import logging

# Import OCR service
from app.services.ocr_service import extract_text_from_image

# Import STT service
from app.services.stt_service import stt_service

logger = logging.getLogger(__name__)

# --- OCR Fallback Service ---
def _ocr_page_image(page_image: bytes) -> Tuple[str, float]:
    """
    Extract text from a page image using OCR.
    
    This function now uses the real OCR service. If OCR is unavailable or disabled,
    it will gracefully fall back to placeholder text without breaking the pipeline.

    Args:
        page_image (bytes): The image content of a single page.

    Returns:
        A tuple containing the extracted text and a confidence score (0.0 to 1.0).
    """
    logger.info("OCR fallback triggered for a page image")
    return extract_text_from_image(page_image, preprocess=True)

# --- PDF Extraction Service ---
def _extract_pdf(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extracts text and metadata from a PDF file, page by page.
    It attempts to extract digital text first and uses an OCR fallback for pages
    that appear to be scanned images.

    Args:
        file_path (str): The path to the PDF file.

    Returns:
        A tuple containing a list of content blocks (one per page) and the document type
        (e.g., 'pdf-digital' or 'pdf-ocr').
    """
    blocks = []
    doc_type = "pdf-digital"  # Assume digital by default
    extraction_method = "PyMuPDF"
    try:
        doc = fitz.open(file_path)
        # Iterate through each page of the PDF
        for page_num, page in enumerate(doc):
            # Extract text directly. `sort=True` helps maintain reading order.
            page_text = page.get_text("text", sort=True)
            ocr_confidence = 1.0  # High confidence for direct digital extraction

            # OCR Fallback Check: If a page has very little selectable text, it might be a scanned image.
            if len(page_text.strip()) < 50:  # Arbitrary threshold for text density
                # Confirm there are images on the page before triggering OCR
                if page.get_images():
                    doc_type = "pdf-ocr"
                    extraction_method = "PyMuPDF_OCR_fallback"
                    # Convert the page to an image (pixmap) for OCR processing
                    pix = page.get_pixmap()
                    page_text, ocr_confidence = _ocr_page_image(pix.tobytes("png"))

            # Store the content and metadata for the page
            blocks.append({
                "text": page_text,
                "metadata": {
                    "source_filename": os.path.basename(file_path),
                    "page_number": page_num + 1,
                    "extraction_method": extraction_method,
                    "ocr_confidence": ocr_confidence
                }
            })
        doc.close()
    except Exception as e:
        print(f"Error processing PDF {file_path}: {e}")
        return [], "pdf-error"
    return blocks, doc_type

# --- HTML Extraction Service ---
def _extract_html(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extracts the main article content from an HTML file, stripping out boilerplate
    like navigation, ads, and footers using the 'trafilatura' library.

    Args:
        file_path (str): The path to the HTML file.

    Returns:
        A list containing a single block of the main extracted content and the doc type 'html'.
    """
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        html_content = f.read()
    
    soup = BeautifulSoup(html_content, 'html.parser')
    title = soup.title.string if soup.title else "N/A"

    # Use trafilatura to extract the core article text, keeping tables.
    main_content = trafilatura_extract(html_content, include_comments=False, include_tables=True)
    blocks = [{
        "text": main_content or "",
        "metadata": {
            "source_filename": os.path.basename(file_path),
            "title": title,
            "extraction_method": "trafilatura"
        }
    }]
    return blocks, "html"

# --- DOCX Extraction Service ---
def _extract_docx(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extracts text from a .docx file, grouping paragraphs by headings to preserve context.
    
    Instead of treating each paragraph as a separate block, this function accumulates 
    body text under the preceding heading. This ensures that sections like 
    "1. General Policy" and its following content are kept together for the chunker.

    Args:
        file_path (str): The path to the .docx file.

    Returns:
        A list of grouped blocks and the doc type 'docx'.
    """
    blocks = []
    try:
        doc = docx.Document(file_path)
        
        current_text_parts = []
        current_metadata = {
            "source_filename": os.path.basename(file_path),
            "start_paragraph_index": 0,
            "style": "body",
            "extraction_method": "python-docx-grouped"
        }
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
                
            style_name = para.style.name
            is_heading = style_name.startswith('Heading')
            
            # If we hit a new heading and we have accumulated text, save the current block
            if is_heading and current_text_parts:
                # Save previous block
                full_text = "\n".join(current_text_parts)
                blocks.append({
                    "text": full_text,
                    "metadata": current_metadata
                })
                
                # Reset for new section
                current_text_parts = []
                # Determine new style level
                level = style_name.split(' ')[-1] if ' ' in style_name else '1'
                current_metadata = {
                    "source_filename": os.path.basename(file_path),
                    "start_paragraph_index": i,
                    "style": f"h{level}",
                    "extraction_method": "python-docx-grouped"
                }
            
            # Add current paragraph (heading or body) to the accumulator
            current_text_parts.append(text)
            
            # If this was the first paragraph of the doc (and maybe not a heading), set metadata
            if not blocks and not current_text_parts[:-1]:
                 if not is_heading:
                     current_metadata["style"] = "body"
        
        # Append the final block if content remains
        if current_text_parts:
            full_text = "\n".join(current_text_parts)
            blocks.append({
                "text": full_text,
                "metadata": current_metadata
            })

        # Basic table extraction (kept separate for now as tables are distinct structures)
        for table_idx, table in enumerate(doc.tables):
            for row_idx, row in enumerate(table.rows):
                row_cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(filter(None, row_cells))
                if row_text:
                    blocks.append({
                        "text": f"Table Row: {row_text}",
                        "metadata": {
                            "source_filename": os.path.basename(file_path),
                            "style": "table_row",
                            "table_index": table_idx,
                            "row_index": row_idx,
                            "extraction_method": "python-docx"
                        }
                    })

    except Exception as e:
        print(f"Error processing DOCX {file_path}: {e}")
        return [], "docx-error"
    return blocks, "docx"

# --- Spreadsheet Extraction Service ---
def _process_dataframe(df: pd.DataFrame, file_path: str, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Helper to process a dataframe into grouped chunks with header injection.
    Groups 20 rows per chunk to maintain context for RAG while including column names.
    """
    blocks = []
    headers = ", ".join(df.columns.astype(str).tolist())
    rows_per_chunk = 20
    
    # Replace pandas NaN with None for consistent handling
    df = df.where(pd.notna(df), None)
    
    num_rows = len(df)
    for i in range(0, num_rows, rows_per_chunk):
        chunk_df = df.iloc[i : i + rows_per_chunk]
        row_strings = []
        
        for idx, row in chunk_df.iterrows():
            # Create string representation, skipping None values
            row_items = [f"{col}: {val}" for col, val in row.items() if val is not None]
            if row_items:
                # Store the 1-based row index (Header is row 1, so data starts at 2)
                row_strings.append(f"Row {idx + 2}: " + ", ".join(row_items))
        
        if row_strings:
            # Inject headers at the top of every chunk
            combined_text = f"Table Columns: {headers}\n---\n" + "\n".join(row_strings)
            
            metadata = {
                "source_filename": os.path.basename(file_path),
                "row_range": f"{i + 2}-{min(i + rows_per_chunk + 1, num_rows + 1)}",
                "extraction_method": "pandas_grouped_with_headers"
            }
            if sheet_name:
                metadata["sheet_name"] = sheet_name
                
            blocks.append({
                "text": combined_text,
                "metadata": metadata
            })
    return blocks

def _extract_spreadsheet(file_path: str, file_type: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extracts data from spreadsheet files (XLSX, CSV), grouping rows and injecting headers.
    
    Args:
        file_path (str): The path to the spreadsheet file.
        file_type (str): The type of spreadsheet ('xlsx' or 'csv').

    Returns:
        A list of blocks, where each block contains a group of rows with header context.
    """
    all_blocks = []
    try:
        if file_type == 'xlsx':
            xls = pd.ExcelFile(file_path)
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name).dropna(how='all')
                if not df.empty:
                    sheet_blocks = _process_dataframe(df, file_path, sheet_name)
                    all_blocks.extend(sheet_blocks)
        else:  # For CSV
            # Attempt to read CSV with different encodings
            encodings = ['utf-8', 'cp1252', 'latin-1']
            df = None
            
            for encoding in encodings:
                try:
                    df = pd.read_csv(file_path, encoding=encoding).dropna(how='all')
                    logger.info(f"Successfully read CSV '{file_path}' with encoding '{encoding}'")
                    break
                except UnicodeDecodeError:
                    continue
                except Exception as e:
                    logger.warning(f"Failed to read CSV with {encoding}: {e}")
                    continue
            
            if df is None:
                raise ValueError(f"Could not read CSV file {file_path} with any of the supported encodings ({encodings})")

            if not df.empty:
                csv_blocks = _process_dataframe(df, file_path)
                all_blocks.extend(csv_blocks)
                
    except Exception as e:
        logger.error(f"Error processing spreadsheet {file_path}: {str(e)}")
        return [], f"{file_type}-error"
        
    return all_blocks, file_type

# --- SQL Dump Extraction Service (Handles both INSERT and COPY) ---
def _extract_sql(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extracts data from SQL dump files, supporting both standard INSERT statements
    and PostgreSQL's 'COPY ... FROM stdin' format.

    Args:
        file_path (str): The path to the .sql file.

    Returns:
        A list of blocks, where each block is a row of data as a JSON string.
    """
    blocks = []
    print(f"--- Starting Unified SQL Extraction for {file_path} ---")
    try:
        # Try to detect encoding first
        encodings_to_try = ['utf-8', 'utf-16', 'utf-16le', 'utf-16be', 'latin-1', 'cp1252']
        content = None
        
        for encoding in encodings_to_try:
            try:
                with open(file_path, 'r', encoding=encoding, errors='ignore') as f:
                    test_content = f.read()
                # Check if content looks reasonable (no excessive null bytes)
                if test_content and test_content.count('\x00') / len(test_content) < 0.1:
                    content = test_content
                    print(f"Successfully read SQL file with encoding: {encoding}")
                    break
            except:
                continue
        
        if not content:
            # Fallback to utf-8 with error handling
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()

        # --- Process COPY statements using line-by-line parsing ---
        print("\n--- Processing for COPY statements ---")
        found_copy_blocks = 0
        lines = content.split('\n')
        in_copy_block = False
        current_table = None
        current_columns = []
        
        for i, line in enumerate(lines):
            line = line.strip()
            
            # Check for start of COPY block
            if line.startswith('COPY ') and 'FROM stdin;' in line:
                copy_match = re.match(r'COPY\s+([\w\.]+)\s*\((.*?)\)\s+FROM\s+stdin;', line, re.IGNORECASE)
                if copy_match:
                    found_copy_blocks += 1
                    current_table = copy_match.group(1).strip()
                    current_columns = [col.strip().strip('"') for col in copy_match.group(2).split(',')]
                    in_copy_block = True
                    print(f"Found COPY block for table '{current_table}' with columns: {current_columns}")
                    
            # Check for end of COPY block
            elif in_copy_block and line == '\\.':
                in_copy_block = False
                current_table = None
                current_columns = []
                
            # Process data lines within COPY block
            elif in_copy_block and line:
                values = line.split('\t')
                if len(current_columns) == len(values):
                    row_data = {col: (val if val != '\\N' else None) for col, val in zip(current_columns, values)}
                    blocks.append({
                        "text": json.dumps(row_data),
                        "metadata": { 
                            "source_filename": os.path.basename(file_path), 
                            "table_name": current_table, 
                            "extraction_method": "psql_copy" 
                        }
                    })
                else:
                    print(f"[WARNING] Column count mismatch in table '{current_table}': expected {len(current_columns)}, got {len(values)}. Skipping row.")
        
        if found_copy_blocks == 0:
            print("No COPY blocks found.")

        # --- Process INSERT statements using sqlparse ---
        print("\n--- Processing for INSERT statements ---")
        found_insert_blocks = 0
        for stmt in sqlparse.parse(content):
            if stmt.get_type() != 'INSERT':
                continue
            
            found_insert_blocks += 1
            table_name = "unknown"
            into_seen = False
            for token in stmt.tokens:
                if into_seen and isinstance(token, sqlparse.sql.Identifier):
                    table_name = token.get_real_name()
                    break
                if token.is_keyword and token.normalized == 'INTO':
                    into_seen = True

            columns_part = next((t for t in stmt.tokens if isinstance(t, sqlparse.sql.Parenthesis)), None)
            column_names = []
            if columns_part:
                id_list = next((t for t in columns_part.tokens if isinstance(t, sqlparse.sql.IdentifierList)), None)
                if id_list:
                    column_names = [col.get_real_name() for col in id_list.get_identifiers()]

            values_part = next((t for t in stmt.tokens if isinstance(t, sqlparse.sql.Values)), None)
            if not values_part:
                continue

            for row_parens in values_part.get_sublists():
                if isinstance(row_parens, sqlparse.sql.Parenthesis):
                    id_list = next((t for t in row_parens.tokens if isinstance(t, sqlparse.sql.IdentifierList)), None)
                    if not id_list: continue
                    
                    row_values = [identifier.normalized.strip("'\"") for identifier in id_list.get_identifiers()]
                    
                    if column_names and len(column_names) == len(row_values):
                        row_data = dict(zip(column_names, row_values))
                    else:
                        row_data = {f"col_{j}": val for j, val in enumerate(row_values)}
                    
                    blocks.append({
                        "text": json.dumps(row_data),
                        "metadata": { "source_filename": os.path.basename(file_path), "table_name": table_name, "extraction_method": "sql_insert" }
                    })
        
        if found_insert_blocks > 0:
            print(f"Found and processed {found_insert_blocks} INSERT statements.")
        else:
            print("No INSERT statements found.")

        print(f"--- SQL Extraction Complete. Found {len(blocks)} total blocks. ---")

    except Exception as e:
        print(f"Error processing SQL file {file_path}: {e}")
        import traceback
        traceback.print_exc()
        return [], "sql-error"
    
    return blocks, "sql"


# --- JSONL Extraction Service ---
def _extract_jsonl(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extracts data from a JSONL (JSON Lines) file, where each line is a separate JSON object.

    Args:
        file_path (str): The path to the .jsonl file.

    Returns:
        A list of blocks, where each block is one line (a JSON object) from the file.
    """
    blocks = []
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            for i, line in enumerate(f):
                if line.strip():
                    blocks.append({
                        "text": line.strip(), 
                        "metadata": {
                            "source_filename": os.path.basename(file_path),
                            "record_number": i + 1,
                            "extraction_method": "line-by-line"
                        }
                    })
    except Exception as e:
        print(f"Error processing JSONL {file_path}: {e}")
        return [], "jsonl-error"
    return blocks, "jsonl"

# --- Plain Text & Markdown Extraction Service ---
def _extract_text(file_path: str, file_type: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    A generic extractor for plain text files (.txt, .md, etc.). It reads the entire
    file content into a single block.

    Args:
        file_path (str): The path to the text file.
        file_type (str): The type of file ('txt', 'md', etc.).

    Returns:
        A list containing a single block with the file's content.
    """
    encodings = ['utf-8', 'cp1252', 'latin-1']
    content = ""
    success = False
    
    for encoding in encodings:
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                content = f.read()
                success = True
                logger.info(f"Successfully read text file '{file_path}' with encoding '{encoding}'")
                break
        except UnicodeDecodeError:
            continue
        except Exception as e:
            logger.warning(f"Error reading text file with {encoding}: {e}")
            continue
    
    if not success:
        # Final fallback with errors='ignore'
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
            
    blocks = [{
        "text": content,
        "metadata": {
            "source_filename": os.path.basename(file_path),
            "extraction_method": "plain-text"
        }
    }]
    return blocks, file_type

# --- Audio Extraction Service ---
async def _extract_audio(file_path: str, file_type: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extracts text from audio files using Speech-to-Text (STT) service.
    
    Args:
        file_path (str): The path to the audio file.
        file_type (str): The type of audio file ('mp3', 'wav', etc.).
    
    Returns:
        A list containing a single block with the transcribed text and metadata.
    """
    try:
        logger.info(f"Starting audio transcription for {file_path}")
        
        # Use STT service to transcribe audio
        transcription_result = await stt_service.transcribe_audio(file_path)
        
        if not transcription_result["text"].strip():
            logger.warning(f"No speech detected in audio file: {file_path}")
            return [], f"{file_type}-no-speech"
        
        # Create a single block with transcribed content
        blocks = [{
            "text": transcription_result["text"],
            "metadata": {
                "source_filename": os.path.basename(file_path),
                "original_format": file_type,
                "transcription_provider": transcription_result["provider"],
                "language": transcription_result["language"],
                "confidence": transcription_result["confidence"],
                "extraction_method": "stt_transcription",
                "content_type": "audio_transcription"
            }
        }]
        
        logger.info(f"Audio transcription completed for {file_path}. Provider: {transcription_result['provider']}")
        return blocks, f"{file_type}-transcribed"
        
    except Exception as e:
        logger.error(f"Error processing audio file {file_path}: {e}")
        return [], f"{file_type}-error"


# --- File Extractor Router ---
# This dictionary maps file extensions to their corresponding extraction functions.
# This acts as a router to select the correct strategy based on file type.
FILE_EXTRACTORS = {
    'pdf': _extract_pdf,
    'html': _extract_html,
    'htm': _extract_html,
    'docx': _extract_docx,
    'xlsx': lambda p: _extract_spreadsheet(p, 'xlsx'),
    'csv': lambda p: _extract_spreadsheet(p, 'csv'),
    'sql': _extract_sql,
    'jsonl': _extract_jsonl,
    'txt': lambda p: _extract_text(p, 'txt'),
    'md': lambda p: _extract_text(p, 'md'),
    # Audio file types
    'wav': lambda p: _extract_audio(p, 'wav'),
    'mp3': lambda p: _extract_audio(p, 'mp3'),
    'mp4': lambda p: _extract_audio(p, 'mp4'),
    'm4a': lambda p: _extract_audio(p, 'm4a'),
    'aac': lambda p: _extract_audio(p, 'aac'),
    'ogg': lambda p: _extract_audio(p, 'ogg'),
    'webm': lambda p: _extract_audio(p, 'webm'),
    'flac': lambda p: _extract_audio(p, 'flac'),
}

# --- Main Entry Point ---
async def extract_data_from_file(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Identifies the file type and routes it to the appropriate extraction function.
    Now supports async extraction for audio files.

    Args:
        file_path (str): The path to the file to be processed.

    Returns:
        A tuple containing the list of extracted data blocks and the determined document type.
        If the file type is not supported, it falls back to a plain text extraction.
    """
    # Get the file extension to determine the file type.
    _, file_extension = os.path.splitext(file_path)
    file_type = file_extension.lower().strip('.')

    # Define audio file types that require async processing
    audio_types = {'wav', 'mp3', 'mp4', 'm4a', 'aac', 'ogg', 'webm', 'flac'}
    
    # Look up the appropriate extractor function from the router dictionary.
    extractor = FILE_EXTRACTORS.get(file_type)
    
    if extractor:
        # If an extractor is found, call it (async for audio, sync for others)
        if file_type in audio_types:
            blocks, doc_type = await extractor(file_path)
        else:
            blocks, doc_type = await anyio.to_thread.run_sync(extractor, file_path)
        
        for block in blocks:
            block["metadata"]["doc_type"] = doc_type
        return blocks, doc_type
    else:
        # If the file type is unknown, use the plain text extractor as a fallback.
        print(f"No specific extractor for file type '{file_type}', using plain text fallback.")
        blocks, doc_type = await anyio.to_thread.run_sync(_extract_text, file_path, 'unknown')
        for block in blocks:
            block["metadata"]["doc_type"] = doc_type
        return blocks, doc_type

# Backward compatibility: sync wrapper for non-audio files
def extract_data_from_file_sync(file_path: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Synchronous wrapper for extract_data_from_file for backward compatibility.
    Only use this for non-audio files.
    """
    import asyncio
    
    # Get the file extension to determine the file type.
    _, file_extension = os.path.splitext(file_path)
    file_type = file_extension.lower().strip('.')
    
    # Define audio file types that require async processing
    audio_types = {'wav', 'mp3', 'mp4', 'm4a', 'aac', 'ogg', 'webm', 'flac'}
    
    if file_type in audio_types:
        # For audio files, we need to run the async function
        return asyncio.run(extract_data_from_file(file_path))
    else:
        # For non-audio files, we can call the sync version directly
        extractor = FILE_EXTRACTORS.get(file_type)
        
        if extractor:
            blocks, doc_type = extractor(file_path)
            for block in blocks:
                block["metadata"]["doc_type"] = doc_type
            return blocks, doc_type
        else:
            print(f"No specific extractor for file type '{file_type}', using plain text fallback.")
            blocks, doc_type = _extract_text(file_path, 'unknown')
            for block in blocks:
                block["metadata"]["doc_type"] = doc_type
            return blocks, doc_type
